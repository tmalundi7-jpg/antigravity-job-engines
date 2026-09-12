def _safe_str(v):
    if v is None:
        return ''
    s = str(v)
    # Fix common encoding artifacts
    s = s.replace('Aâ¬', '£').replace('Â£', '£').replace('A£', '£')
    return s

import logging
from collections import defaultdict

def _ensure_openpyxl():
    import openpyxl
    return openpyxl

def _header_style(openpyxl, fill_hex):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    thin = Side(style="thin", color="AAAAAA")
    return dict(
        font=Font(bold=True, color="FFFFFF", size=11),
        fill=PatternFill("solid", fgColor=fill_hex),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=Border(left=thin, right=thin, top=thin, bottom=thin),
    )

def _style_cell(cell, **kwargs):
    for k, v in kwargs.items(): setattr(cell, k, v)

def _auto_width(ws, min_w=10, max_w=60):
    for col in ws.columns:
        ltr = col[0].column_letter
        mx = min_w
        for cell in col:
            try: mx = max(mx, len(str(cell.value or "")))
            except Exception: pass
        ws.column_dimensions[ltr].width = min(mx + 2, max_w)

def export_executive_report_excel(state: dict, output_path: str):
    openpyxl = _ensure_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side
    wb = Workbook()
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    matches = state.get("ranked_matches", [])
    ws2 = wb.active
    ws2.title = "Ranked Matches"
    h2 = ["Rank","Score /100","Job Title","Grade/Band","Employer","Location","Salary",
          "Vector Similarity","Reasoning","Source URL"]
    for c, h in enumerate(h2, 1):
        _style_cell(ws2.cell(row=1, column=c, value=_safe_str(h) if isinstance(h, str) else h), **_header_style(openpyxl, "1F4E79"))
    ws2.row_dimensions[1].height = 30
    for r, m in enumerate(matches, 2):
        vals = [r-1, m.get("final_score",0), m.get("job_title",""), m.get("grade_or_band",""),
                m.get("company",""), m.get("location",""),
                m.get("salary_range","Competitive"), m.get("similarity_score",""),
                m.get("reasoning",""), m.get("source_url","")]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(row=r, column=c, value=_safe_str(v) if isinstance(v, str) else v)
            cell.border = border
    _auto_width(ws2)
    wb.save(output_path)
    return output_path

def export_available_roles_excel(state: dict, output_path: str):
    openpyxl = _ensure_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Border, Side
    wb = Workbook()
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    parsed_jobs = [j for j in state.get("parsed_jobs", []) if not j.get("is_generic")]
    companies = state.get("companies", [])
    company_lookup = {c.get("name","").lower(): c for c in companies}

    enriched = []
    for job in parsed_jobs:
        co = job.get("company", "Unknown")
        meta = company_lookup.get(co.lower(), {})
        enriched.append({
            "title":       job.get("title","Unknown Role"),
            "grade":       job.get("grade_or_band","N/A"),
            "employer":    co,
            "sector":      meta.get("sector","Public Sector"),
            "location":    job.get("location","UK"),
            "salary":      job.get("salary_range", job.get("salary","Competitive")),
            "source_url":  job.get("source_url",""),
        })

    all_headers = ["No.","Job Title","Grade/Band","Employer","Sector","Location","Salary","Source URL"]

    def write_sheet(ws, jobs, fill_hex, hdr_hex):
        for c, h in enumerate(all_headers, 1):
            _style_cell(ws.cell(row=1,column=c,value=_safe_str(h) if isinstance(h, str) else h), **_header_style(openpyxl,hdr_hex))
        ws.row_dimensions[1].height = 30
        for r, job in enumerate(jobs, 2):
            is_portal = job.get("source_url", "").startswith("http") and job.get("salary", "") == "See website"
            if is_portal:
                fill = "FFF2CC"  # Yellow highlight for portal-only rows
            else:
                fill = fill_hex if r%2==0 else "FFFFFF"
            vals = [r-1, job["title"], job["grade"], job["employer"], job["sector"],
                    job["location"], job["salary"], job["source_url"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=_safe_str(v) if isinstance(v, str) else v)
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.border = border
                if is_portal and c == 2:  # italicise the title on portal rows
                    from openpyxl.styles import Font as _Font
                    cell.font = _Font(italic=True, color="7F6000")
        ws.freeze_panes = "A2"
        if jobs: ws.auto_filter.ref = f"A1:H{len(jobs)+1}"
        _auto_width(ws)

    # 1. All Roles
    ws1 = wb.active; ws1.title = "All Public Roles"
    write_sheet(ws1, enriched, "E2EFDA", "375623")
    
    # 2. NHS
    nhs_jobs = [j for j in enriched if "NHS" in j["sector"]]
    if nhs_jobs: write_sheet(wb.create_sheet("NHS Roles"), nhs_jobs, "EBF3FB", "1F4E79")
        
    # 3. Civil Service
    cs_jobs = [j for j in enriched if "Civil Service" in j["sector"]]
    if cs_jobs: write_sheet(wb.create_sheet("Civil Service Roles"), cs_jobs, "FFF9E6", "BF8F00")

    # 4. Local Government / Regulators
    other_jobs = [j for j in enriched if "NHS" not in j["sector"] and "Civil Service" not in j["sector"]]
    if other_jobs: write_sheet(wb.create_sheet("Other Public Bodies"), other_jobs, "FCE4D6", "C65911")

    wb.save(output_path)
    return output_path


def export_generic_roles_excel(state: dict, output_path: str):
    openpyxl = _ensure_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Border, Side
    wb = Workbook()
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    parsed_jobs = state.get("parsed_jobs", [])
    generic_jobs = [j for j in parsed_jobs if j.get("is_generic")]
    
    companies = state.get("companies", [])
    company_lookup = {c.get("name","").lower(): c for c in companies}

    enriched = []
    for job in generic_jobs:
        co = job.get("company", "Unknown")
        meta = company_lookup.get(co.lower(), {})
        enriched.append({
            "title":       job.get("title","Unknown Role"),
            "grade":       job.get("grade_or_band","N/A"),
            "employer":    co,
            "sector":      meta.get("sector","Public Sector"),
            "location":    job.get("location","UK"),
            "salary":      job.get("salary_range", job.get("salary","Competitive")),
            "source_url":  job.get("source_url",""),
        })

    all_headers = ["No.","Job Title","Grade/Band","Employer","Sector","Location","Salary","Source URL"]

    def write_sheet(ws, jobs, fill_hex, hdr_hex):
        for c, h in enumerate(all_headers, 1):
            _style_cell(ws.cell(row=1,column=c,value=_safe_str(h) if isinstance(h, str) else h), **_header_style(openpyxl,hdr_hex))
        ws.row_dimensions[1].height = 30
        for r, job in enumerate(jobs, 2):
            is_portal = job.get("source_url", "").startswith("http") and job.get("salary", "") == "See website"
            if is_portal:
                fill = "FFF2CC"  # Yellow highlight for portal-only rows
            else:
                fill = fill_hex if r%2==0 else "FFFFFF"
            vals = [r-1, job["title"], job["grade"], job["employer"], job["sector"],
                    job["location"], job["salary"], job["source_url"]]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=_safe_str(v) if isinstance(v, str) else v)
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.border = border
                if is_portal and c == 2:  # italicise the title on portal rows
                    from openpyxl.styles import Font as _Font
                    cell.font = _Font(italic=True, color="7F6000")
        ws.freeze_panes = "A2"
        if jobs: ws.auto_filter.ref = f"A1:H{len(jobs)+1}"
        _auto_width(ws)

    ws1 = wb.active; ws1.title = "Generic Public Roles"
    write_sheet(ws1, enriched, "E2EFDA", "375623")
    
    wb.save(output_path)
    return output_path
